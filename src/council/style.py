"""Read-only PPTX style measurement with content-minimizing derived reports."""

from __future__ import annotations

import hashlib
import json
import statistics
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

import yaml
from pptx import Presentation
from pptx.dml.color import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

from council.store import CouncilError, safe_path

START = "<!-- council:observations:start -->"
END = "<!-- council:observations:end -->"
LIMITATIONS = [
    "Direct run formatting omits inherited fonts and font sizes; missing values are counted.",
    "Backgrounds distinguish direct RGB, theme references, and inherited/other fills.",
    "Text categories use lexical cues; they do not establish whether results are completed.",
    "Caption and logo counts are candidates, not semantic identifications.",
    "Embedded raster plots, equations in images, and plot-label readability require manual review.",
    "Pictures without embedded data are counted separately; linked resources are never fetched.",
    "Speaker notes, slide text, images, filenames, and source paths are not exported.",
    "Layout names and theme fonts are observed formatting metadata; review before disclosure.",
    "Counts use decks, slides, or runs as labeled; related decks can bias corpus-level summaries.",
]


def category(text: str) -> str:
    """Classify structural cues without exporting original text."""
    lower = text.lower()
    for label, cues in (
        ("closing", ("thank", "questions", "conclusion", "summary")),
        ("planned", ("future", "next step", "planned", "will ")),
        ("preliminary", ("preliminary", "ongoing", "in progress")),
        ("completed", ("completed", "validated", "verified")),
        ("methods", ("method", "approach", "algorithm")),
        ("results", ("results", "comparison", "convergence")),
        ("opening", ("outline", "agenda", "introduction", "objective")),
    ):
        if any(cue in lower for cue in cues):
            return label
    return "unclassified"


def quantiles(values: list[float | int]) -> dict:
    """Summarize distributions deterministically, including empty samples."""
    if not values:
        return {"samples": 0, "min": None, "median": None, "max": None}
    return {
        "samples": len(values),
        "min": min(values),
        "median": statistics.median(values),
        "max": max(values),
    }


def shapes_recursive(shapes):
    """Include grouped shapes in observations."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shapes_recursive(shape.shapes)


def inspect_deck(path: Path) -> dict:
    """Extract observable formatting without persisting confidential slide text."""
    with path.open("rb") as stream:
        digest = hashlib.file_digest(stream, "sha256").hexdigest()
    deck = Presentation(path)
    fonts: Counter = Counter()
    sizes: Counter = Counter()
    backgrounds: Counter = Counter()
    layouts: Counter = Counter()
    categories: Counter = Counter()
    title_positions = []
    word_counts = []
    paragraph_lengths = []
    bullet_lengths = []
    title_sizes = []
    body_sizes = []
    image_hashes: Counter = Counter()
    features = Counter(
        pictures=0,
        tables=0,
        charts=0,
        equation_xml=0,
        caption_candidates=0,
        explicit_bullet_paragraphs=0,
        unknown_font_runs=0,
        unknown_size_runs=0,
        pictures_without_embedded_data=0,
    )
    structures = []
    for slide in deck.slides:
        layouts[slide.slide_layout.name] += 1
        fill = slide.background.fill
        if fill.type is not None and fill.type == 1:
            color = fill.fore_color
            if color.type == MSO_COLOR_TYPE.RGB:
                backgrounds[f"rgb:{color.rgb}"] += 1
            elif color.type == MSO_COLOR_TYPE.SCHEME:
                backgrounds[f"theme:{color.theme_color}"] += 1
            else:
                backgrounds["other"] += 1
        else:
            backgrounds["inherited-or-nonsolid"] += 1
        title = slide.shapes.title
        if title is not None:
            title_positions.append(
                {
                    "x": round(title.left / deck.slide_width, 4),
                    "y": round(title.top / deck.slide_height, 4),
                    "width": round(title.width / deck.slide_width, 4),
                    "height": round(title.height / deck.slide_height, 4),
                }
            )
        slide_words = 0
        slide_text = []
        for shape in shapes_recursive(slide.shapes):
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                features["pictures"] += 1
                try:
                    blob = shape.image.blob
                except ValueError:
                    features["pictures_without_embedded_data"] += 1
                else:
                    image_hashes[hashlib.sha256(blob).hexdigest()] += 1
            if shape.has_table:
                features["tables"] += 1
            if shape.has_chart:
                features["charts"] += 1
            features["equation_xml"] += len(shape._element.xpath(".//*[local-name()='oMath']"))
            if not shape.has_text_frame:
                continue
            slide_text.append(shape.text)
            if shape.text.lower().lstrip().startswith(("fig.", "figure ", "table ")):
                features["caption_candidates"] += 1
            for paragraph in shape.text_frame.paragraphs:
                count = len(paragraph.text.split())
                slide_words += count
                if count:
                    paragraph_lengths.append(count)
                bullet = paragraph._p.xpath("./a:pPr/a:buChar | ./a:pPr/a:buAutoNum")
                if bullet:
                    features["explicit_bullet_paragraphs"] += 1
                    bullet_lengths.append(count)
                for run in paragraph.runs:
                    if run.font.name:
                        fonts[run.font.name] += 1
                    else:
                        features["unknown_font_runs"] += 1
                    if run.font.size:
                        size = round(run.font.size.pt, 2)
                        sizes[str(size)] += 1
                        (
                            title_sizes
                            if title is not None and shape.shape_id == title.shape_id
                            else body_sizes
                        ).append(size)
                    else:
                        features["unknown_size_runs"] += 1
        word_counts.append(slide_words)
        structural = category(title.text if title is not None else "")
        structures.append(structural)
        categories[category(" ".join(slide_text))] += 1
    themes = []
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if name.startswith("ppt/theme/theme") and name.endswith(".xml"):
                tree = ET.fromstring(archive.read(name))
                namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                colors = []
                for node in tree.findall(".//a:clrScheme/*/*", namespace):
                    colors.append(node.get("lastClr") or node.get("val"))
                typefaces = sorted(
                    {
                        node.get("typeface")
                        for node in tree.findall(".//a:fontScheme//a:latin", namespace)
                        if node.get("typeface")
                    }
                )
                themes.append({"colors": colors, "latin_typefaces": typefaces})
    return {
        "sha256": digest,
        "bytes": path.stat().st_size,
        "slides": len(deck.slides),
        "dimensions_inches": {
            "width": round(deck.slide_width / 914400, 4),
            "height": round(deck.slide_height / 914400, 4),
        },
        "themes": themes,
        "direct_fonts_by_run": dict(sorted(fonts.items())),
        "direct_font_sizes_pt_by_run": dict(sorted(sizes.items())),
        "font_hierarchy_pt": {"title": quantiles(title_sizes), "body": quantiles(body_sizes)},
        "backgrounds": dict(backgrounds),
        "layouts": dict(layouts),
        "title_positions_normalized": title_positions,
        "words_per_slide": quantiles(word_counts),
        "words_per_paragraph": quantiles(paragraph_lengths),
        "words_per_explicit_bullet": quantiles(bullet_lengths),
        "features": dict(features),
        "repeated_image_candidates": sum(1 for n in image_hashes.values() if n > 1),
        "structural_category_counts": dict(Counter(structures)),
        "repeated_category_pairs": {
            f"{a}->{b}": n
            for (a, b), n in Counter(zip(structures, structures[1:], strict=False)).items()
            if n > 1
        },
        "opening_category": structures[0] if structures else None,
        "closing_category": structures[-1] if structures else None,
        "status_cue_slide_counts": dict(categories),
        "plot_labeling": "manual review required",
    }


def ingest(root: Path, source: Path, dry_run: bool = False) -> str:
    """Measure a folder of decks, then update generated sections without erasing prose."""
    source = safe_path(root, source)
    if not source.is_dir():
        raise CouncilError("Presentation source must be an existing repository-relative directory")
    paths = sorted(p for p in source.rglob("*") if p.is_file() and p.suffix.lower() == ".pptx")
    decks = []
    for path in paths:
        safe_path(root, path.relative_to(root))
        try:
            decks.append(inspect_deck(path))
        except Exception as exc:
            raise CouncilError(
                f"Could not ingest a PPTX ({type(exc).__name__}); no reports were written"
            ) from exc
    decks.sort(key=lambda item: item["sha256"])
    profile = {
        "schema_version": "1.0",
        "deck_count": len(decks),
        "unique_deck_hashes": len({deck["sha256"] for deck in decks}),
        "slide_count": sum(d["slides"] for d in decks),
        "observations": decks,
        "limitations": LIMITATIONS,
        "interpretations": [
            "No intentions inferred. Compare candidate conventions with Santiago before adoption."
        ],
    }
    observations = "# Presentation style observations\n\n"
    observations += f"Direct observations: {len(decks)} decks, {profile['slide_count']} slides.\n\n"
    if not decks:
        observations += "No decks available. Known preferences remain the only style guidance.\n\n"
    else:
        dimensions: Counter = Counter()
        corpus_fonts: Counter = Counter()
        corpus_sizes: Counter = Counter()
        corpus_backgrounds: Counter = Counter()
        openings: Counter = Counter()
        closings: Counter = Counter()
        features: Counter = Counter()
        title_y = []
        density = []
        for deck in decks:
            size = deck["dimensions_inches"]
            dimensions[f"{size['width']} x {size['height']} inches"] += 1
            corpus_fonts.update(deck["direct_fonts_by_run"])
            corpus_sizes.update(deck["direct_font_sizes_pt_by_run"])
            corpus_backgrounds.update(deck["backgrounds"])
            openings[deck["opening_category"] or "empty"] += 1
            closings[deck["closing_category"] or "empty"] += 1
            features.update(deck["features"])
            title_y.extend(p["y"] for p in deck["title_positions_normalized"])
            if deck["words_per_slide"]["median"] is not None:
                density.append(deck["words_per_slide"]["median"])

        def top(counter: Counter) -> str:
            return "; ".join(f"{key}: {count}" for key, count in counter.most_common(3)) or "none"

        observations += (
            f"- Unique file hashes: {profile['unique_deck_hashes']}.\n"
            f"- Most frequent slide dimensions (deck counts): {top(dimensions)}.\n"
            f"- Most frequent directly formatted fonts (run counts): {top(corpus_fonts)}.\n"
            f"- Most frequent explicit font sizes (points; run counts): {top(corpus_sizes)}.\n"
            f"- Background observations (slide counts): {top(corpus_backgrounds)}.\n"
            f"- Median of per-deck median words/slide: {quantiles(density)['median']}.\n"
            f"- Median title top/slide height: {quantiles(title_y)['median']}.\n"
            f"- Pictures: {features['pictures']}; tables: {features['tables']}; "
            f"charts: {features['charts']}; equation XML nodes: {features['equation_xml']}.\n"
            f"- Opening title categories (lexical candidates): {top(openings)}.\n"
            f"- Closing title categories (lexical candidates): {top(closings)}.\n\n"
        )
        observations += (
            "See the profile for per-deck formatting distributions and structural cue counts.\n\n"
        )
    observations += (
        "## Limits of interpretation\n\n" + "\n".join(f"- {item}" for item in LIMITATIONS) + "\n"
    )
    guide = safe_path(root, "style/presentation-style-guide.md")
    original = guide.read_text(encoding="utf-8")
    generated = (
        START
        + "\n\n"
        + observations.replace("# Presentation style observations", "## Corpus observations", 1)
        + "\n"
        + END
    )
    if START in original or END in original:
        if (
            original.count(START) != 1
            or original.count(END) != 1
            or original.index(END) < original.index(START)
        ):
            raise CouncilError(
                "Invalid generated markers in style guide; preserve manual edits and repair markers"
            )
        updated = (
            original[: original.index(START)]
            + generated
            + original[original.index(END) + len(END) :]
        )
    else:
        updated = original.rstrip() + "\n\n" + generated + "\n"
    outputs = {
        "style/derived/deck-inventory.json": json.dumps(
            {
                "schema_version": "1.0",
                "decks": [
                    {k: d[k] for k in ("sha256", "bytes", "slides", "dimensions_inches")}
                    for d in decks
                ],
            },
            indent=2,
        )
        + "\n",
        "style/derived/presentation-style-profile.yaml": yaml.safe_dump(profile, sort_keys=False),
        "style/derived/style-observations.md": observations,
        "style/presentation-style-guide.md": updated,
    }
    # Resolve every destination before writing any output.
    destinations = {safe_path(root, name): text for name, text in outputs.items()}
    if dry_run:
        return json.dumps(
            {"dry_run": True, "deck_count": len(decks), "would_write": list(outputs)}, indent=2
        )
    for path, text in destinations.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8", newline="\n")
    return (
        f"Ingested {len(decks)} decks / {profile['slide_count']} slides. "
        "Review derived reports before disclosure."
    )
